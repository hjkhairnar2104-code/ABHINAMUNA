import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT_DIR = r"d:\EVFINAL"
RESEARCH_IMG_DIR = os.path.join(ROOT_DIR, "frontend", "public", "research")
DOCX_OUTPUT_PATH = os.path.join(ROOT_DIR, "IEEE_Battery_Health_SOH_Prognostics_Paper.docx")

print("Generating Word Document (.docx)...", flush=True)

doc = Document()

# Set standard 0.8 in margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title_p.add_run("Physics-Informed Machine Learning with LightGBM for Lithium-Ion Battery Health Prognostics & Knee-Point Lifecycle Estimation Across Heterogeneous Chemistries\n")
t_run.font.size = Pt(17)
t_run.font.bold = True
t_run.font.name = "Arial"
t_run.font.color.rgb = RGBColor(15, 23, 42)

# Authors
a_p = doc.add_paragraph()
a_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
a_run = a_p.add_run("SmartEV AIML Advanced Energy Storage Systems Research Group\nAdvanced Electric Vehicle Energy Storage and Battery Health Laboratory\nTarget: IEEE Transactions on Transportation Electrification")
a_run.font.size = Pt(10)
a_run.font.italic = True
a_run.font.color.rgb = RGBColor(71, 85, 105)

# Abstract
abs_p = doc.add_paragraph()
abs_t = abs_p.add_run("ABSTRACT\n")
abs_t.font.bold = True
abs_t.font.size = Pt(10)
abs_t.font.color.rgb = RGBColor(30, 58, 138)
abs_run = abs_p.add_run(
    "Accurate State of Health (SOH) prognostics and Remaining Useful Life (RUL) lifecycle forecasting are essential for the operational safety, thermal runaway mitigation, warranty economics, and second-life grid reuse of Lithium-Ion Battery (LIB) systems in Electric Vehicles (EVs). Traditional electrochemical pseudo-two-dimensional (P2D) models suffer from prohibitive computational overhead that prevents real-time Edge-BMS implementation, while purely data-driven black-box neural networks frequently violate thermodynamic irreversibility constraints and fail near rapid non-linear knee-point inflection.\n\n"
    "In this paper, we establish a leakage-controlled, physics-informed prognostic benchmark centered on LightGBM with histogram-based binning. We conduct an exhaustive audit of empirical EV fleet telemetry (N = 4,732 commercial vehicle records across LFP, NMC, NCA, and LMO cells) and high-fidelity laboratory electrochemical cycling data (N = 10,600 cycles across 11 prismatic LCO cells from CALCE). To eliminate circular evaluation, we implement a strict Cell-Separated Group Split ensuring test cells are completely unseen during training.\n\n"
    "On unseen held-out cells, LightGBM achieves an R² score of 0.9880, a Mean Absolute Error (MAE) of 1.395% SOH, and a Root Mean Squared Error (RMSE) of 1.879% SOH, executing single-sample inference in only 0.019 ms (19 µs). Feature ablation confirms that partial charging curve features combined with cumulative operational context achieve parity with full-cycle discharge data. Furthermore, we implement Conformalized Quantile Regression (CQR), attaining an 85.16% empirical coverage on unseen cells at nominal 90% confidence with a tight 6.51% SOH mean interval width, exposing cell-to-cell manufacturing heterogeneity as an inherent physical boundary."
)
abs_run.font.size = Pt(9.5)

# Section 1
doc.add_heading("1. Executive Summary & Forensic Data Audit", level=1)
doc.add_paragraph("Prior academic drafts and baseline models in this project reported inflated prognostic metrics (e.g. R² = 0.9979, MAE = 0.429%) due to random row-level train/test splits that placed adjacent charge cycles of the same physical battery cell in both training and test partitions. In this investigation, we perform a complete forensic data audit, enforce zero-leakage inductive cell partitioning, and establish genuine empirical performance benchmarks.")

t1 = doc.add_table(rows=1, cols=4)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
h = t1.rows[0].cells
h[0].text = "Dataset / Artifact"
h[1].text = "Reported Metric / Claim"
h[2].text = "Verified Reality on Disk"
h[3].text = "Resolution & Corrective Action"

audit_data = [
    ("Consolidated Master Dataset", "399,079 cycles, 362 cells, 6 sources", "10,601 text rows (4.00 MB), 11 cells, 1 source (CALCE LCO); 162.28 MB trailing NUL bytes", "Interrupted disk write truncated CSV. Isolated clean 10,600-cycle CALCE benchmark."),
    ("EV Fleet Telemetry", "Claimed 50,000 commercial records", "Exactly 4,732 records across 4 chemistries (10.57 MB)", "Corrected sample size to N = 4,732; retrained models on true fleet distribution."),
    ("Internal Resistance (CALCE)", "Claimed 73.12% importance in CALCE SOH", "ir_ohm is 100.0% missing (NaN) in all 10,600 CALCE cycle records", "Internal resistance was not measured in CALCE. Excluded ir_ohm to eliminate imputation bias."),
    ("Data Partitioning Protocol", "Random 80/20 train_test_split on cyclic rows", "Severe cell leakage: adjacent cycles of identical cells evaluated in test", "Implemented strict Cell-Separated Group Split (GroupKFold) with zero leakage."),
    ("Unseen Cell Performance", "Claimed R² = 0.9979, MAE = 0.429%", "Genuine unseen-cell score: R² = 0.9880, MAE = 1.395% SOH, RMSE = 1.879%", "Reported authentic, publication-grade generalization without data leakage."),
    ("Uncertainty Calibration (CQR)", "Claimed marginal coverage 0.895 without width", "Nominal 90% CQR yields 85.16% empirical coverage, mean width 6.51% SOH", "Uncovered cell-level group shift (CS2: 76.41% vs CX2: 99.28%); documented physical limits.")
]

for row in audit_data:
    cells = t1.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val

# Section 2
doc.add_heading("2. Benchmark Results on Unseen Battery Cells", level=1)
doc.add_paragraph("To replicate real-world deployment on completely unseen batteries, the 11 physical CALCE cells are partitioned into: Training (6 cells, 5,950 cycles), Validation (2 cells, 2,283 cycles), Calibration (1 cell, 905 cycles), and Held-Out Unseen Test (2 cells: CS2_38 and CX2_36, 1,462 cycles).")

t2 = doc.add_table(rows=1, cols=6)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
h2 = t2.rows[0].cells
h2[0].text = "Model Architecture"
h2[1].text = "RMSE (% SOH)"
h2[2].text = "MAE (% SOH)"
h2[3].text = "R² Score"
h2[4].text = "Train Time"
h2[5].text = "Latency"

bench_data = [
    ("Dummy Regressor (Mean Baseline)", "17.169%", "12.282%", "-0.0031", "0.059 s", "<0.001 ms"),
    ("Ridge Regression (alpha=1.0)", "3.045%", "1.913%", "0.9684", "0.035 s", "0.002 ms"),
    ("Random Forest (50 Trees, Depth 10)", "2.629%", "1.884%", "0.9765", "2.620 s", "0.011 ms"),
    ("Gradient Boosted Trees (GBDT)", "2.340%", "1.702%", "0.9814", "5.131 s", "0.004 ms"),
    ("XGBoost Regressor", "1.995%", "1.370%", "0.9865", "2.415 s", "0.008 ms"),
    ("LightGBM Regressor (Champion)", "1.879%", "1.395%", "0.9880", "0.394 s", "0.019 ms")
]

for row in bench_data:
    cells = t2.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val

# Add Parity Picture
fig1_path = os.path.join(RESEARCH_IMG_DIR, "fig1_unseen_cell_parity.png")
if os.path.exists(fig1_path):
    doc.add_paragraph()
    doc.add_picture(fig1_path, width=Inches(5.8))
    cap1 = doc.add_paragraph("Fig. 1. Actual vs. Predicted State of Health (SOH) on 1,462 cycles from unseen physical cells under LightGBM (R² = 0.9880, MAE = 1.395%, RMSE = 1.879%).")
    cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap1.runs[0].font.size = Pt(8.5)
    cap1.runs[0].font.italic = True

# Section 3
doc.add_heading("3. Feature Ablation Study: Partial Curves vs. Full Cycles", level=1)
doc.add_paragraph("We evaluate three operational feature sets to investigate whether full discharge data are necessary for accurate SOH estimation:")

t3 = doc.add_table(rows=1, cols=4)
t3.alignment = WD_TABLE_ALIGNMENT.CENTER
h3 = t3.rows[0].cells
h3[0].text = "Feature Configuration"
h3[1].text = "Feature Count"
h3[2].text = "RMSE (% SOH)"
h3[3].text = "R² Score"

abl_data = [
    ("Config 1: Partial Curve Only", "9 features", "2.874%", "0.9719"),
    ("Config 2: Partial Curve + Operating Context", "21 features", "1.879%", "0.9880"),
    ("Config 3: Full-Cycle Benchmark (with Discharge)", "27 features", "1.909%", "0.9876")
]
for row in abl_data:
    cells = t3.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val

fig3_path = os.path.join(RESEARCH_IMG_DIR, "fig3_model_benchmark_bars.png")
if os.path.exists(fig3_path):
    doc.add_paragraph()
    doc.add_picture(fig3_path, width=Inches(5.6))
    cap3 = doc.add_paragraph("Fig. 2. Error metric comparison across all candidate algorithms on unseen held-out battery cells.")
    cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap3.runs[0].font.size = Pt(8.5)
    cap3.runs[0].font.italic = True

# Section 4
doc.add_heading("4. Conformalized Quantile Regression (CQR) Uncertainty Quantification", level=1)
doc.add_paragraph("Using pinball loss quantile models calibrated on physical cell CALCE_CS2_37 (N = 905 cycles), nominal 90% prediction intervals were formed using conformal correction factor Q_{1-alpha} = 1.4081% SOH. On held-out unseen cells, empirical coverage reached 85.16% with an average interval width of 6.51% SOH.")

fig2_path = os.path.join(RESEARCH_IMG_DIR, "fig2_cqr_prediction_intervals.png")
if os.path.exists(fig2_path):
    doc.add_paragraph()
    doc.add_picture(fig2_path, width=Inches(5.8))
    cap2 = doc.add_paragraph("Fig. 3. Calibrated nominal 90% CQR prediction interval trajectory over the lifecycle of unseen cell CALCE_CS2_38.")
    cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap2.runs[0].font.size = Pt(8.5)
    cap2.runs[0].font.italic = True

# Section 5
doc.add_heading("5. Production API Deployment & Automated Consistency Tests", level=1)
doc.add_paragraph("All production artifacts in ml-service/app/models were retrained and verified with compress=3. Automated test suite test_deployment_consistency.py verified 9 critical checks with 100% pass rate: model loading, feature schemas, direct offline-to-API parity (97.0% SOH), extreme value handling, and CQR interval responses.")

doc.save(DOCX_OUTPUT_PATH)
print(f"SUCCESS: Generated {DOCX_OUTPUT_PATH} ({os.path.getsize(DOCX_OUTPUT_PATH):,} bytes)", flush=True)
